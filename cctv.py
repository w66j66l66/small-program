#库
import requests
import re
from bs4 import BeautifulSoup
import os
import concurrent.futures
from concurrent.futures import ThreadPoolExecutor, as_completed

def parse_cctv_article(html):
    soup = BeautifulSoup(html, 'lxml')

    title_tag = soup.find('h1')
    title = title_tag.get_text(strip=True) if title_tag else (soup.title.string.strip() if soup.title else '')

    pub = None
    for meta_name in ('pubdate', 'publishdate', 'article:published_time', 'ptime'):
        m = soup.find('meta', attrs={'name': meta_name}) or soup.find('meta', attrs={'property': meta_name})
        if m and m.get('content'):
            pub = m['content']
            break
    if not pub:
        time_tag = soup.find(class_=re.compile(r'(time|date|pub|ptime)', re.I))
        if time_tag:
            pub = time_tag.get_text(strip=True)

    candidates = [
        ('div', {'class': 'cnt_bd'}),
        ('div', {'class': 'content'}),
        ('div', {'id': 'content'}),
        ('div', {'class': re.compile(r'(article|content|main)', re.I)}),
        ('div', {'id': re.compile(r'(article|content|main)', re.I)}),
        ('article', {}),
    ]
    article_text = ''
    for name, attrs in candidates:
        node = soup.find(name, attrs=attrs)
        if node:
            parts = [p.get_text(strip=True) for p in node.find_all(['p', 'div']) if p.get_text(strip=True)]
            article_text = '\n\n'.join(parts).strip()
            if len(article_text) > 50:
                break

    if not article_text:
        paragraphs = [p.get_text(strip=True) for p in soup.find_all('p') if p.get_text(strip=True)]
        article_text = '\n\n'.join(paragraphs).strip()

    return {
        'title': title,
        'published': pub,
        'content': article_text
    }

def save_articles_pdf(articles, out_path):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception:
        raise RuntimeError('需要安装 reportlab：pip install reportlab')

    possible_fonts = [
        r"C:\Windows\Fonts\msyh.ttf",
        r"C:\Windows\Fonts\msjh.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simfang.ttf"
    ]
    font_path = None
    for p in possible_fonts:
        if os.path.exists(p):
            font_path = p
            break
    if not font_path:
        raise RuntimeError('未找到中文字体文件，请将系统字体路径写入 possible_fonts 列表或安装中文字体。')

    font_name = 'UserChinese'
    pdfmetrics.registerFont(TTFont(font_name, font_path))

    doc = SimpleDocTemplate(out_path, pagesize=A4,
                            rightMargin=40, leftMargin=40, topMargin=60, bottomMargin=60)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MyTitle', fontName=font_name, fontSize=16, leading=20, spaceAfter=6))
    styles.add(ParagraphStyle(name='MyMeta', fontName=font_name, fontSize=10, leading=12, textColor=colors.grey))
    styles.add(ParagraphStyle(name='MyBody', fontName=font_name, fontSize=12, leading=16))

    story = []
    if not articles:
        story.append(Paragraph('未抓取到任何文章。', styles['MyBody']))
    else:
        for idx, art in enumerate(articles, 1):
            title = art.get('title') or ('未命名 第{}篇'.format(idx))
            pub = art.get('published') or ''
            content = art.get('content') or ''
            content = content.replace('&', '&amp;')
            content = content.replace('\r\n', '\n').replace('\r', '\n')
            content = content.replace('\n\n', '<br/><br/>').replace('\n', '<br/>')
            story.append(Paragraph(title, styles['MyTitle']))
            if pub:
                story.append(Paragraph(pub, styles['MyMeta']))
            story.append(Spacer(1, 6))
            story.append(Paragraph(content, styles['MyBody']))
            if idx != len(articles):
                story.append(PageBreak())

    doc.build(story)

def save_articles_docx(articles, out_path):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.oxml.ns import qn
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    except Exception:
        raise RuntimeError('需要安装 python-docx：pip install python-docx')

    # 
    preferred_fonts = ['微软雅黑', 'Microsoft YaHei', 'SimHei', '黑体', '宋体', 'SimSun']
    font_name = preferred_fonts[0]

    doc = Document()
    # 主要用于正文
    style = doc.styles['Normal']
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    style.font.size = Pt(12)

    if not articles:
        doc.add_paragraph('未抓取到任何文章。')
    else:
        for idx, art in enumerate(articles, 1):
            title = art.get('title') or f'未命名 第{idx}篇'
            pub = art.get('published') or ''
            content = art.get('content') or ''

            # 标题
            p = doc.add_paragraph()
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

            # 发布时间
            if pub:
                p2 = doc.add_paragraph()
                run2 = p2.add_run(pub)
                run2.font.size = Pt(10)
                run2.font.name = font_name
                run2._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

            doc.add_paragraph()  # 空行

           
            content = content.replace('\r\n', '\n').replace('\r', '\n')
            
            paragraphs = []
            buffer = []
            for line in content.split('\n'):
                line = line.strip()
                if not line:
                    if buffer:
                        paragraphs.append(' '.join(buffer))
                        buffer = []
                else:
                    buffer.append(line)
            if buffer:
                paragraphs.append(' '.join(buffer))

            for para in paragraphs:
                p_body = doc.add_paragraph()
                run_body = p_body.add_run(para)
                run_body.font.size = Pt(12)
                run_body.font.name = font_name
                run_body._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

            
            if idx != len(articles):
                doc.add_page_break()

    doc.save(out_path)

if __name__ == '__main__':
    url = 'https://news.cctv.com/2019/07/gaiban/cmsdatainterface/page/news_1.jsonp?cb=news'
    try:
        resp = requests.get(url, timeout=10)
        text = resp.text
    except Exception as e:
        print('获取列表失败：', e)
        matches = []
    else:
        pattern = r'"url":"(https://news\.cctv\.com/\d{4}/\d{2}/\d{2}/[^"]+)"'
        matches = re.findall(pattern, text)

    if not matches:
        print('未找到任何文章链接。')
    else:
        articles = []

        def fetch_parse(u):
            try:
                session = requests.Session()
                r = session.get(u, timeout=10)
                r.encoding = 'utf-8'
                return parse_cctv_article(r.text)
            except Exception as e:
                print('抓取出错：', u, e)
                return None

        max_workers = min(16, max(1, len(matches)))
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_url = {executor.submit(fetch_parse, u): u for u in matches}
            completed = 0
            for fut in as_completed(future_to_url):
                completed += 1
                url_done = future_to_url[fut]
                try:
                    data = fut.result()
                    if data:
                        articles.append({
                            'url': url_done,
                            'title': data.get('title'),
                            'published': data.get('published'),
                            'content': data.get('content')
                        })
                        print(f'[{completed}/{len(matches)}] 标题:', data.get('title'))
                except Exception as e:
                    print('任务异常：', url_done, e)

        script_dir = os.path.dirname(os.path.abspath(__file__))
        out_docx = os.path.join(script_dir, '每日新闻.docx')
        try:
            save_articles_docx(articles, out_docx)
            print('已生成 Word 文档：', out_docx)
        except Exception as e:
            print('生成 Word 失败：', e)



